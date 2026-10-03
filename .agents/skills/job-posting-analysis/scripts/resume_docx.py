"""Build or conservatively patch a resume from a Codex-authored, evidence-linked plan."""
import argparse
import json
import re
import platform
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def paragraphs(doc):
    yield from doc.paragraphs
    def cells(tables):
        seen = set()
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen: continue
                    seen.add(id(cell._tc))
                    yield from cell.paragraphs
                    yield from cells(cell.tables)
    yield from cells(doc.tables)


def validate(plan, source):
    if not plan.get('company') or not plan.get('position'):
        raise ValueError('company and position are required')
    if plan.get('jd_status') != 'COMPLETE':
        raise ValueError('A complete JD is required for resume generation')
    if not plan.get('source_resume') or not plan.get('jd_url'):
        raise ValueError('source_resume and jd_url are required')
    mode = plan.get('mode', 'rebuild')
    if mode not in ('rebuild', 'patch'): raise ValueError('mode must be rebuild or patch')
    entries = plan.get('edits', []) if mode == 'patch' else [x for sec in plan.get('sections', []) for x in sec.get('items', [])]
    if not entries: raise ValueError('No resume content')
    for entry in entries:
        text = entry.get('after') if mode == 'patch' else entry.get('text')
        if not isinstance(text, str) or not text.strip(): raise ValueError('Empty resume item')
        if not entry.get('source_ref') or not entry.get('reason'): raise ValueError('Evidence reference and reason required')
        quotes = entry.get('evidence', [])
        if not quotes or not all(isinstance(q, str) and q.strip() and q in source for q in quotes):
            raise ValueError('Evidence must be exact excerpts present in the source text')
        # Reject newly introduced numeric claims. This is not a semantic fact checker.
        tokens = set(re.findall(r'\d+(?:[.,]\d+)*', text))
        original = set(re.findall(r'\d+(?:[.,]\d+)*', '\n'.join(quotes)))
        if tokens - original: raise ValueError('New numeric claim without supporting evidence: '+str(tokens-original))
        if mode == 'patch' and not entry.get('before'): raise ValueError('Patch requires exact before text')
    if mode == 'rebuild':
        if not plan.get('candidate_name'): raise ValueError('candidate_name required')
        if plan['candidate_name'] not in source: raise ValueError('Name not present in source')
        for line in plan.get('contact', []):
            if line not in source: raise ValueError('Contact information not present in source')
    return entries


def build(plan, source_text, output_dir):
    entries = validate(plan, source_text)
    out = Path(output_dir).resolve(); out.mkdir(parents=True, exist_ok=True)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', plan['company']+'_'+plan['position'])[:120]
    destination = out/(name+'_이력서.docx')
    changes_path = out/(name+'_변경내역.md')
    if destination.exists() or changes_path.exists(): raise FileExistsError('Output exists; choose a new output directory')
    original = Path(plan['source_resume']).resolve()
    if destination == original: raise ValueError('Cannot overwrite the original resume')
    notes = []
    if plan.get('mode', 'rebuild') == 'patch':
        if original.suffix.lower() != '.docx': raise ValueError('Patch mode requires a DOCX source')
        doc = Document(original)
        available = list(paragraphs(doc)); assignments = []; used = set()
        for entry in entries:
            matches = [p for p in available if p.text == entry['before']]
            if len(matches) != 1: raise ValueError('before text must match exactly one paragraph')
            p = matches[0]
            if id(p._p) in used: raise ValueError('Duplicate patch target')
            used.add(id(p._p))
            # Fields, drawings and hyperlinks need specialized edits; do not silently delete them.
            if p._p.xpath('.//w:hyperlink | .//w:drawing | .//w:fldChar'):
                raise ValueError('Complex paragraph cannot be patched safely')
            assignments.append((p,entry))
        for p,entry in assignments:
            if p.runs:
                p.runs[0].text = entry['after']
                for run in p.runs[1:]: run.text = ''
            else: p.add_run(entry['after'])
        notes.append('DOCX 원본의 문단·표·페이지 설정을 유지했습니다. 수정 문단 내부의 혼합 글자 서식은 첫 run 서식으로 통일됩니다. 문장 길이가 바뀌면 페이지 나눔은 달라질 수 있습니다.')
    else:
        doc = Document()
        section = doc.sections[0]
        section.page_width=Inches(8.5); section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=Inches(.7)
        section.left_margin=section.right_margin=Inches(.75)
        font_name = plan.get('font') or ('Apple SD Gothic Neo' if platform.system() == 'Darwin' else 'Malgun Gothic' if platform.system() == 'Windows' else 'Noto Sans CJK KR')
        for style in doc.styles:
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        for style_name, size in [('Normal',11),('Title',20),('Heading 1',13)]:
            style = doc.styles[style_name]
            style.font.name=font_name; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),font_name)
        doc.styles['Normal'].paragraph_format.space_after=Pt(6)
        doc.add_paragraph(plan['candidate_name'], 'Title')
        for line in plan.get('contact', []): doc.add_paragraph(line)
        for sec in plan['sections']:
            doc.add_heading(sec['heading'], level=1)
            for item in sec['items']:
                p=doc.add_paragraph(item['text'], style='List Bullet' if item.get('bullet') else 'Normal')
                p.paragraph_format.widow_control=True
        notes.append('원문 내용을 재구성한 DOCX입니다. 원본 PDF·이미지의 디자인을 복제하지 않습니다. 섹션 순서는 입력 계획을 따릅니다.')
    doc.core_properties.author=''; doc.core_properties.last_modified_by=''
    doc.save(destination)
    # Structural read-back, not visual layout verification.
    saved=Document(destination); text='\n'.join(p.text for p in paragraphs(saved))
    for entry in entries:
        expected=entry.get('after') if plan.get('mode')=='patch' else entry['text']
        if expected not in text: raise ValueError('DOCX read-back missing generated content')
    lines=['# 이력서 변경 내역','',f"대상 회사: {plan['company']}",f"대상 직무: {plan['position']}",f"JD: {plan['jd_url']}",f"원본: {original.name}",'',*notes,'','문구의 사실 일치와 원문 누락 여부는 Codex 검토가 필요합니다. 숫자·근거 검사는 의미 수준의 검증을 대신하지 않습니다.','시각적 레이아웃 검증은 별도 필요합니다.','']
    for i,entry in enumerate(entries,1):
        lines.extend([f'## 변경 {i}',f"이전: {entry.get('before', '재구성')}" , f"수정: {entry.get('after', entry.get('text'))}",f"이유: {entry['reason']}",f"근거 위치: {entry['source_ref']}", '근거 원문: '+' / '.join(entry['evidence']),''])
    lines.extend(['## 추가 확인 사항',*['- '+n for n in plan.get('review_notes', [])]])
    changes_path.write_text('\n'.join(lines),encoding='utf-8')
    return {'docx':str(destination),'changes':str(changes_path),'visual_qa':'not_checked'}

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--plan',required=True);ap.add_argument('--source-text',required=True);ap.add_argument('--output-dir',required=True);args=ap.parse_args()
    print(json.dumps(build(json.loads(Path(args.plan).read_text(encoding='utf-8')),Path(args.source_text).read_text(encoding='utf-8'),args.output_dir),ensure_ascii=False))
