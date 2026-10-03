import importlib.util, tempfile, unittest
from pathlib import Path
from docx import Document
p=Path(__file__).resolve().parents[1]/'.agents/skills/job-posting-analysis/scripts/resume_docx.py';s=importlib.util.spec_from_file_location('resume_docx',p);m=importlib.util.module_from_spec(s);s.loader.exec_module(m)
class ResumeTests(unittest.TestCase):
 def plan(self):
  return {'company':'회사','position':'PM','jd_url':'https://example.com/job','jd_status':'COMPLETE','source_resume':'resume.pdf','candidate_name':'테스트 지원자','sections':[{'heading':'경력','items':[{'text':'LLM 평가 기준 설계','evidence':['모델 평가 기준 설계'],'source_ref':'1쪽 경력','reason':'JD 표현에 맞춰 정리'}]}]}
 def test_rebuild(self):
  with tempfile.TemporaryDirectory() as t:
   r=m.build(self.plan(),'테스트 지원자\n모델 평가 기준 설계',t)
   self.assertTrue(Path(r['docx']).exists());self.assertTrue(Path(r['changes']).exists())
   with self.assertRaises(FileExistsError):m.build(self.plan(),'테스트 지원자\n모델 평가 기준 설계',t)
 def test_unsupported_evidence(self):
  with self.assertRaises(ValueError):m.validate(self.plan(),'테스트 지원자')
 def test_invented_number(self):
  p=self.plan();p['sections'][0]['items'][0]['text']='평가 정확도 99% 개선'
  with self.assertRaises(ValueError):m.validate(p,'테스트 지원자\n모델 평가 기준 설계')
 def test_incomplete_jd(self):
  p=self.plan();p['jd_status']='PARTIAL'
  with self.assertRaises(ValueError):m.validate(p,'테스트 지원자\n모델 평가 기준 설계')
 def test_patch_preserves_original(self):
  with tempfile.TemporaryDirectory() as t:
   original=Path(t)/'original.docx';doc=Document();doc.add_paragraph('모델 평가 기준 설계');doc.save(original);before=original.read_bytes()
   p=self.plan();p.update(mode='patch',source_resume=str(original),edits=[{'before':'모델 평가 기준 설계','after':'LLM 평가 기준 설계','evidence':['모델 평가 기준 설계'],'source_ref':'1쪽','reason':'회사 언어 반영'}])
   r=m.build(p,'모델 평가 기준 설계',Path(t)/'out');self.assertEqual(original.read_bytes(),before);self.assertEqual(Document(r['docx']).paragraphs[0].text,'LLM 평가 기준 설계')
if __name__=='__main__':unittest.main()
